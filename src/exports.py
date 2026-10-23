"""Downloads are built inside one job, including spreadsheet formula protection."""
from __future__ import annotations
import csv
import io
import json
import re
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import PatternFill

COLUMNS = ['gene_or_target','disease','expression_direction','disease_effect','relation_type',
           'species','evidence_type','evidence_sentence','paper_id','title','source_location','file_name','page','chunk_id',
           'source_url','quote_verified','needs_review','notes','statistical_evidence','synthetic',
           'evidence_id','review_status','review_revision','reviewer','review_reason','reviewed_at',
           'machine_result','content_level','source_type','source_quality_flags','section','paragraph','pmcid','doi',
           'abstract_sha256','fulltext_sha256','fulltext_license','relation_scope','population_context','measurement_type',
           'support_sentence','measurement_sentence','support_status','support_flags','support_messages',
           'support_quote_verified','measurement_quote_verified','measurement_hint','relationship_validated','support_check_version',
           'task_profile','comparison_type','disease_group','normal_group','comparison_sentence','comparison_source_location','comparison_page',
           'significance','baseline_status','baseline_flags','baseline_messages','comparison_quote_verified','baseline_check_version']


def safe_cell(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', value)
    if isinstance(value, str) and value.lstrip().startswith(('=', '+', '-', '@')):
        return "'" + value
    return value


def export_bytes(rows: list[dict], extension: str) -> bytes:
    if extension == 'json':
        return json.dumps(rows, ensure_ascii=False, indent=2).encode('utf-8')
    if extension == 'csv':
        buffer = io.StringIO(newline='')
        writer = csv.writer(buffer)
        writer.writerow(COLUMNS)
        for row in rows:
            writer.writerow([safe_cell(row.get(c, '')) for c in COLUMNS])
        return buffer.getvalue().encode('utf-8-sig')
    if extension != 'xlsx':
        raise ValueError('不支持的导出格式。')
    book = Workbook()
    sheet = book.active
    sheet.title = '证据'
    sheet.append(COLUMNS)
    for row in rows:
        sheet.append([safe_cell(row.get(c, '')) for c in COLUMNS])
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='DCE9E7')
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions['H'].width = 85
    sheet.column_dimensions['J'].width = 45
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def write_exports(directory: Path, rows: list[dict]):
    directory.mkdir(parents=True, exist_ok=True)
    for extension in ['json','csv','xlsx']:
        (directory/f'evidence.{extension}').write_bytes(export_bytes(rows, extension))
