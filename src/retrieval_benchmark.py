"""Offline ablation of the shared candidate pipeline; no new ranking algorithm.

Judgments belong to a fixed query AND exact supplied passage. Missing judgments
and insufficient information never become negative relevance labels.
"""
from __future__ import annotations

import copy
import hashlib
import json
import statistics
import time
from collections import Counter
from pathlib import Path

METHODS = {
    'bm25': ('关键词评分（BM25＋命中加分）', False, False),
    'hybrid': ('关键词与向量融合', True, False),
    'hybrid_rerank': ('融合加重排序', True, True),
}
DECISIONS = ['有本任务证据', '仅主题相关，没有本任务证据', '不相关', '信息不足，暂不判断']
HEADERS = ['核对编号', '目标基因', '目标疾病', '论文编号', '来源与位置', '标题',
           '所给原文（仅按此判断）', '人工判断', '支持原句（有证据时复制）', '复核人', '备注']


def digest(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
        separators=(',', ':')).encode()).hexdigest()


def save_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def _checked_bundle(bundle):
    if bundle.get('schema') != 'frozen-retrieval-comparison-v1':
        raise ValueError('不是支持的冻结评测材料格式。')
    queries = bundle['queries']; chunks = bundle['local_index']['chunks']
    if not 1 <= len(queries) <= 20 or len(chunks) > 2000:
        raise ValueError('本工具限定1至20组问题、最多2000个本地片段。')
    if len({q['query_id'] for q in queries}) != len(queries):
        raise ValueError('问题编号重复。')
    for q in queries:
        if not q.get('gene') or not q.get('disease') or len(q.get('online_chunks', [])) > 100:
            raise ValueError('查询目标为空或在线冻结候选超过100条。')
        for c in [*chunks, *q.get('online_chunks', [])]:
            if not c.get('chunk_id') or not isinstance(c.get('context'), str):
                raise ValueError('候选缺少编号或原文。')


def compare(bundle, output, *, ks=(5, 10, 15), repeats=3, coarse=500, encode=500):
    """Call canonical stages with identical pools, aliases, quotas and context.

Cold preparation is not included in warm medians. Measured order rotates across
repetitions. Final k stays in the canonical call because its reranking pool is
itself a function of k; a single top-15 run cannot stand in for all budgets.
    """
    _checked_bundle(bundle)
    ks = tuple(ks)
    if (not ks or len(set(ks)) != len(ks) or any(type(k) is not int or not 1 <= k <= 50 for k in ks)
            or not 1 <= repeats <= 5 or not max(ks) <= encode <= 2000 or not 1 <= coarse <= 2000):
        raise ValueError('最终数量、重复次数或候选预算无效。')
    output = Path(output)
    output.mkdir(parents=True, exist_ok=False)
    save_json(output/'status.json', {'state':'running', 'completed_queries':0})
    import_started=time.monotonic()
    from retrieval import retrieve_candidates, score_chunk_bm25
    from retrieval_pipeline import rank_candidates
    core_import_seconds=time.monotonic()-import_started
    runs = []; preparations = []; references = []; query_reports = []
    for query_index, q in enumerate(bundle['queries']):
        print(f"[{query_index+1}/{len(bundle['queries'])}] {q['query_id']} {q['gene']} · {q['disease']}", flush=True)
        index = bundle['local_index']; ga = q.get('gene_aliases', []); da = q.get('disease_aliases', [])
        started = time.monotonic()
        local = retrieve_candidates(index, q['gene'], q['disease'], top_k=coarse,
            gene_aliases=ga, disease_aliases=da, require_gene=True)
        online = [c for item in q.get('online_chunks', []) if (c := score_chunk_bm25(
            item, index, q['gene'], q['disease'], ga, da, True)) is not None]
        merged = sorted([*local, *online], key=lambda c:c.get('retrieval_score', 0), reverse=True)
        # Canonical rank_candidates also validates conflicting duplicate IDs.
        unique = {}
        for c in merged:
            if c['chunk_id'] in unique and unique[c['chunk_id']]['context'] != c['context']:
                raise ValueError('同一候选编号对应不同原文。')
            unique.setdefault(c['chunk_id'], c)
        pool = list(unique.values())
        if len(pool) > encode:
            raise ValueError('冻结候选超过共同向量预算；请明确缩小材料或提高预算，不能静默让三方法使用不同池。')
        coarse_seconds = time.monotonic()-started
        signature = digest(pool)
        for c in pool:
            references.append(dict(query_id=q['query_id'], gene=q['gene'], disease=q['disease'],
                review_id=digest([q['query_id'],c['chunk_id'],c['context']])[:20], **copy.deepcopy(c)))
        common = dict(index=index, encode_top_n=encode, gene_aliases=ga,
            disease_aliases=da, per_paper_limit=0)

        def run_one(method, k, repeat, phase):
            _, hybrid, reranker = METHODS[method]
            began = time.monotonic()
            selected, trace = rank_candidates(copy.deepcopy(pool), q['gene'], q['disease'],
                top_k=k, use_hybrid=hybrid, use_reranker=reranker, **common)
            elapsed = time.monotonic()-began
            expected = 'bm25' + ('+dense' if hybrid else '') + ('+cross_encoder' if reranker and len(pool)>1 else '')
            if trace['actual_method'] != expected:
                raise RuntimeError('请求的方法发生降级，停止评测：'+str(trace['actual_method']))
            ids = [c['chunk_id'] for c in selected]
            if len(ids) != len(set(ids)) or not set(ids) <= set(unique) or len(ids)>k:
                raise RuntimeError('评测返回重复、越界或超预算片段。')
            if digest(pool) != signature:
                raise RuntimeError('某方法修改了其他方法共用的候选输入。')
            return dict(query_id=q['query_id'], method=method, k=k, repeat=repeat, phase=phase,
                input_sha256=signature, candidate_count=len(pool), elapsed_seconds=elapsed,
                selected_ids=ids, paper_ids=[c.get('paper_id',c['chunk_id']) for c in selected],
                source_counts=dict(Counter(c.get('source_type','unknown') for c in selected)),
                characters=sum(len(c['context']) for c in selected), trace=trace)

        for method in METHODS:
            preparations.append(run_one(method, max(ks), None, 'preparation'))
        tasks = [(method,k) for k in ks for method in METHODS]
        for repeat in range(repeats):
            offset=(repeat*(len(METHODS)+1)+query_index) % len(tasks)
            for method,k in tasks[offset:]+tasks[:offset]:
                record=run_one(method,k,repeat,'warm')
                runs.append(record)
                with (output/'runs.jsonl').open('a',encoding='utf-8') as f:
                    f.write(json.dumps(record,ensure_ascii=False)+'\n')
        query_reports.append(dict(query_id=q['query_id'],gene=q['gene'],disease=q['disease'],
            gene_aliases=ga,disease_aliases=da,pool_size=len(pool),input_sha256=signature,
            local_candidates=len(local),online_candidates=len(online),coarse_seconds=coarse_seconds))
        save_json(output/'preparations.json', preparations)
        save_json(output/'status.json', {'state':'running','completed_queries':query_index+1})
    aggregate=[]
    for method in METHODS:
        for k in ks:
            selected=[r for r in runs if r['method']==method and r['k']==k]
            aggregate.append(dict(method=method,k=k,measured_runs=len(selected),
                median_warm_rank_seconds=statistics.median(r['elapsed_seconds'] for r in selected),
                median_characters=statistics.median(r['characters'] for r in selected),
                median_unique_papers=statistics.median(len(set(r['paper_ids'])) for r in selected)))
    stable=all(len({tuple(r['selected_ids']) for r in runs if r['query_id']==q['query_id'] and
        r['method']==m and r['k']==k})==1 for q in query_reports for m in METHODS for k in ks)
    manifest=dict(schema='retrieval-review-v1',dataset_sha256=digest(bundle),queries=query_reports,passages=references)
    save_json(output/'review-manifest.json',manifest)
    report=dict(schema='retrieval-comparison-v1',dataset_sha256=digest(bundle),
        sampling=bundle.get('sampling',{}),queries=query_reports,aggregate=aggregate,
        stable_rankings=stable,repeats=repeats,ks=ks,coarse_limit=coarse,encode_limit=encode,
        core_import_seconds=core_import_seconds,llm_calls=0,external_requests=0,quality_metrics=None,
        quality_status='pending_current_task_human_judgments',
        baseline='Canonical BM25 plus its existing gene/disease/expression/relation hit bonuses; not a new plain-BM25 implementation.',
        scope='Frozen historical candidates only; not independent full-corpus dense retrieval, full-corpus recall, end-to-end medical accuracy, or unbiased general latency.',
        timings='Warm medians include candidate-cache checks, query encoding, fusion and optional reranking. Shared coarse search, imports and first preparation are separate. No live online retrieval or LLM.')
    save_json(output/'report.json',report)
    export_review(manifest,output/'中文盲核对表.xlsx')
    export_review(manifest,output/'先核对三组.xlsx',query_ids=[q['query_id'] for q in query_reports[:3]])
    save_json(output/'status.json',{'state':'completed','completed_queries':len(query_reports)})
    return report


def source_values(row):
    return [row['review_id'],row['gene'],row['disease'],row.get('paper_id',''),
        str(row.get('source_type',''))+' · '+str(row.get('source_location','')),
        row.get('title',''),row['context']]


def export_review(manifest, path, *, query_ids=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.utils import get_column_letter
    path=Path(path)
    if path.exists():raise FileExistsError('不能覆盖人工表格。')
    rows=[r for r in manifest['passages'] if query_ids is None or r['query_id'] in query_ids]
    if any(len(r['context'])>32767 for r in rows):raise ValueError('原文超过表格容量，不能静默截断。')
    book=Workbook();guide=book.active;guide.title='先看说明'
    for item in [
        ('目标','仅判断所给文字能否支持：目标基因在疾病状态相对正常对照的表达量差异。人体或动物病理模型均可；不同组织和测量不能混为一谈。'),
        ('填写','打开“只填这页”，只填黄色四列：人工判断、支持原句、复核人、备注。不要改前七列。疾病名称使用英文；type 2 diabetes=2型糖尿病，rheumatoid arthritis=类风湿关节炎，metabolic dysfunction-associated steatohepatitis=代谢功能障碍相关脂肪性肝炎。'),
        ('有本任务证据','目标基因、目标疾病、疾病与正常对照、表达量结果均能从所给文字确认。升高、降低或未见显著差异均可。支持原句连续复制，不要翻译、拼接；需要多句时复制连续一段。'),
        ('仅主题相关','提到基因和疾病，但只有背景、遗传变异、风险、治疗前后、药物或基因干预，未给本任务的表达比较。'),
        ('不相关','所给文字与目标基因或疾病不相关。'),
        ('信息不足','所给文字不足以确认是否有本任务证据，或无法理解比较对象；不强迫填否定。'),
        ('材料范围','只看本表原文。摘要不混入全文知识；本地片段不推断未提供段落。没有结果不表示基因与疾病无关。'),
        ('盲化','未显示哪个方法选中、排名、模型得分或机器答案。按问题分组，组内按不含排名信息的固定编号打散；这只是方法盲化，不是独立专家评审。'),
        ('评测范围','全部为开发材料；既有检索缓存有选择偏差。完整标注后，召回只对这个有限候选池成立，不能声称全库召回。'),
        ('部分填写','可先填前三组的小表。留空或信息不足不当负例；汇总时明确未完成范围。'),
    ]:guide.append(item)
    guide.column_dimensions['A'].width=18;guide.column_dimensions['B'].width=110
    sheet=book.create_sheet('只填这页');sheet.append(HEADERS)
    for r in sorted(rows,key=lambda r:(r['query_id'],digest(r['review_id']))):
        sheet.append(source_values(r)+['','','',''])
    valid=DataValidation(type='list',formula1='"'+','.join(DECISIONS)+'"',allow_blank=True)
    sheet.add_data_validation(valid);valid.add(f'H2:H{max(2,len(rows)+1)}')
    sheet.freeze_panes='H2';sheet.auto_filter.ref=sheet.dimensions
    widths=[24,15,34,22,35,55,105,36,75,18,45]
    for i,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(i)].width=width
    for row in sheet:
        for cell in row:
            cell.alignment=Alignment(wrap_text=True,vertical='top')
            if isinstance(cell.value,str):cell.data_type='s'
            cell.fill=PatternFill('solid',fgColor='FFF2CC' if cell.column>=8 else 'EDF2F7')
        sheet.row_dimensions[row[0].row].height=170 if row[0].row>1 else 35
    for cell in sheet[1]:cell.font=Font(bold=True)
    for row in guide:
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
        guide.row_dimensions[row[0].row].height=65
    book.save(path)


def read_judgments(manifest, workbook):
    from openpyxl import load_workbook
    book=load_workbook(workbook,data_only=False)
    if '只填这页' not in book:raise ValueError('缺少只填这页工作表。')
    sheet=book['只填这页']
    if [c.value for c in sheet[1]]!=HEADERS:raise ValueError('人工表的列名或顺序被修改。')
    expected={r['review_id']:r for r in manifest['passages']};seen=set();labels={}
    for cells in sheet.iter_rows(min_row=2):
        if not any(c.value is not None for c in cells):continue
        if any(c.data_type=='f' for c in cells):raise ValueError('核对表不能包含公式。')
        values=[c.value for c in cells];rid=values[0]
        if rid not in expected or rid in seen:raise ValueError('核对编号无效或重复。')
        seen.add(rid)
        if [v or None for v in values[:7]] != [v or None for v in source_values(expected[rid])]:
            raise ValueError('目标、原文或来源被修改，不能计分。')
        decision,quote,reviewer,note=values[7:]
        if decision in (None,''):continue
        if decision not in DECISIONS:raise ValueError('人工判断不是有效选项。')
        if not str(reviewer or '').strip():raise ValueError('提交判断需要复核人。')
        if decision==DECISIONS[0] and (not isinstance(quote,str) or len(quote.strip())<12 or quote not in expected[rid]['context']):
            raise ValueError('有证据时需要连续复制所给原文中的支持文字。')
        labels[rid]=None if decision==DECISIONS[-1] else int(decision==DECISIONS[0])
    return labels


def score(manifest, runs, labels):
    """P@k only on fully judged outputs; recall only on fully judged fixed pools."""
    by_query={q['query_id']:{} for q in manifest['queries']}
    for row in manifest['passages']:
        by_query[row['query_id']][row['chunk_id']]=labels.get(row['review_id'])
    records=[]
    for qid,pool in by_query.items():
        for method in METHODS:
            for k in sorted({r['k'] for r in runs if r['query_id']==qid and r['method']==method}):
                matched=[r for r in runs if r['query_id']==qid and r['method']==method and r['k']==k]
                if len({tuple(r['selected_ids']) for r in matched})!=1:
                    raise ValueError('重复运行排名不一致，不能只挑一次结果计分。')
                ids=matched[0]['selected_ids']
                if len(ids)!=len(set(ids)) or not set(ids)<=set(pool):raise ValueError('计分输入包含重复或池外编号。')
                known=all(pool[cid] is not None for cid in ids)
                complete=all(v is not None for v in pool.values())
                positives=sum(v==1 for v in pool.values())
                found=sum(pool[cid]==1 for cid in ids)
                records.append(dict(query_id=qid,method=method,k=k,returned=len(ids),
                    judged_returned=sum(pool[cid] is not None for cid in ids),
                    judged_pool=sum(v is not None for v in pool.values()),pool_size=len(pool),
                    precision_at_k=found/k if known else None,
                    precision_returned=found/len(ids) if known and ids else None,
                    recall_within_frozen_pool=found/positives if complete and positives else None,
                    recall_scope='Only the fully judged frozen candidate pool; never full-corpus recall.'))
    return dict(records=records,judged=sum(v is not None for v in labels.values()),
        insufficient=sum(v is None for v in labels.values()),
        unfilled=len(manifest['passages'])-len(labels),
        scope='Development-only evidence relevance in supplied text; no gene-direction extraction accuracy or full-corpus recall.')
