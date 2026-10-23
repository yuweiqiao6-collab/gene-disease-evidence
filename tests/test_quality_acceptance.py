import copy
import hashlib
import json
from pathlib import Path
import pytest
from openpyxl import load_workbook
from conftest import wait_for
from corpus_store import CorpusStore
from jobs import JobManager,Query
from web_settings import WebSettings
from schemas import ExtractionResult
from test_local_library import chunk
from test_research_scripts import script


def reference_case():
    text='IL6 expression was elevated in rheumatoid arthritis patients.'
    return dict(case_id='sample',paper_id='PMID123',gene='IL6',disease='rheumatoid arthritis',
        split='test',title='Synthetic test fixture',source_url='https://example.invalid/123',
        context=text,context_sha256=hashlib.sha256(text.encode()).hexdigest(),
        gold_status='unreviewed',gold_relevant=None,gold_relations=None,reviewer='')


def test_chinese_reference_blank_is_not_gold_and_frozen_cases_are_unchanged(tmp_path):
    module=script('chinese_reference');case=reference_case();original=copy.deepcopy(case)
    workbook=tmp_path/'reference.xlsx';module.export_chinese([case],workbook)
    book=load_workbook(workbook)
    assert book['1论文判断'].freeze_panes=='C2' and book['1论文判断'].column_dimensions['G'].hidden
    assert book['3原文与来源']['F2'].value==case['context']
    output=tmp_path/'labels.jsonl'
    assert module.import_chinese([case],workbook,output)==0
    assert json.loads(output.read_text())['gold_status']=='unreviewed'
    assert case==original
    with pytest.raises(ValueError):module.export_chinese([case],workbook)
    book['3原文与来源']['F2']='Altered visible source';book.save(workbook)
    with pytest.raises(ValueError,match='原文'):module.import_chinese([case],workbook,tmp_path/'altered.jsonl')


def test_chinese_reference_seven_fields_and_strict_source_validation(tmp_path):
    module=script('chinese_reference');case=reference_case();workbook=tmp_path/'ref.xlsx'
    module.export_chinese([case],workbook);book=load_workbook(workbook)
    sheet=book['1论文判断'];sheet['H2']='已完成';sheet['I2']='有';sheet['J2']='否';sheet['K2']='human-reviewer'
    relation=book['2关系答案']
    for col,value in {'D':case['context'],'E':'升高','F':'仅表达差异或标志物','G':'相关性','H':'人',
                      'K':'直接支持目标关系','L':'疾病研究人群','M':'不明确'}.items():relation[f'{col}2']=value
    book.save(workbook);output=tmp_path/'positive.jsonl'
    assert module.import_chinese([case],workbook,output)==1
    result=json.loads(output.read_text());r=result['gold_relations'][0]
    assert r['expression_direction']=='increased' and r['population_context']=='disease_cohort'
    assert r['measurement_type']=='unclear' and r['evidence_sentence']==case['context']
    sheet['G2']='Rewritten source';book.save(workbook)
    with pytest.raises(ValueError,match='原文'):module.import_chinese([case],workbook,tmp_path/'bad.jsonl')


def test_chinese_reference_requires_explicit_negative_and_complete_fields(tmp_path):
    module=script('chinese_reference');case=reference_case();path=tmp_path/'ref.xlsx'
    module.export_chinese([case],path);book=load_workbook(path);sheet=book['1论文判断']
    sheet['H2']='已完成';sheet['I2']='没有';sheet['K2']='reviewer';book.save(path)
    with pytest.raises(ValueError,match='明确确认'):module.import_chinese([case],path,tmp_path/'incomplete.jsonl')
    sheet['J2']='是';book.save(path)
    assert module.import_chinese([case],path,tmp_path/'negative.jsonl')==1
    assert json.loads((tmp_path/'negative.jsonl').read_text())['gold_relations']==[]
    sheet['I2']='有';sheet['J2']='否';book['2关系答案']['D2']=case['context'];book.save(path)
    with pytest.raises(ValueError):module.import_chinese([case],path,tmp_path/'missing-fields.jsonl')


def test_chinese_reference_refuses_formulas_and_text_truncation(tmp_path):
    module=script('chinese_reference');case=reference_case();path=tmp_path/'ref.xlsx'
    module.export_chinese([case],path);book=load_workbook(path);book['1论文判断']['L2']='=1+1';book.save(path)
    with pytest.raises(ValueError,match='公式'):module.import_chinese([case],path,tmp_path/'formula.jsonl')
    case['context']='x'*32768;case['context_sha256']=hashlib.sha256(case['context'].encode()).hexdigest()
    with pytest.raises(ValueError,match='截断'):module.export_chinese([case],tmp_path/'large.xlsx')


def test_selected_paper_coverage_and_no_retry_are_recorded(tmp_path,monkeypatch,owner):
    import jobs
    monkeypatch.setenv('RAG_MODEL_ENV_FILE','');monkeypatch.setenv('LLM_API_KEY','unit-key');monkeypatch.setenv('LLM_MODEL','unit-model')
    corpus=tmp_path/'corpus.sqlite';CorpusStore.build(corpus,[chunk('A',1),chunk('B',1)])
    retry_values=[]
    class Client:
        failure_count=0
        request_count=1
        def __init__(self,**kwargs):pass
        def extract_batch(self,gene,disease,chunks,**kwargs):
            retry_values.append(kwargs['retry'])
            return [(c,[ExtractionResult(evidence_sentence=c['context'],expression_direction='increased')]) for c in chunks]
    monkeypatch.setattr(jobs,'LLMClient',Client)
    service=JobManager(WebSettings(runtime_dir=tmp_path/'runtime',corpus_db=corpus,enable_live=True,access_password='unit-password'))
    try:
        task=service.submit(owner,Query(('IL6',),'rheumatoid arthritis',source='corpus',action='search'))
        found=wait_for(service,owner,task)
        selected=tuple(d['document_id'] for d in found['documents'])
        task=service.submit(owner,Query(('IL6',),'rheumatoid arthritis',source='library',document_ids=selected,
            top_k=1,consent=True,model_retries=0),authorized=True)
        result=wait_for(service,owner,task);assert result['state']=='completed'
        assert retry_values==[0] and result['provenance']['request_limit']==1
        assert sorted(c['candidate_count'] for c in result['provenance']['selected_document_coverage'])==[0,1]
        with pytest.raises(ValueError,match='重试'):
            service.submit(owner,Query(('IL6',),'rheumatoid arthritis',source='library',document_ids=selected,model_retries=2))
    finally:service.close()


def local_plan(tmp_path,monkeypatch):
    module=script('validate_local_workflow')
    monkeypatch.setattr(module,'get_llm_settings',lambda:dict(api_key='unit-key',model='unit-model',base_url='https://example.invalid/v1'))
    corpus=tmp_path/'corpus.sqlite'
    CorpusStore.build(corpus,[chunk(gene,1,f'{gene} expression was elevated in {disease} patients.') for gene,disease in module.TARGETS])
    module.prepare(corpus,tmp_path/'plan')
    return module,corpus,tmp_path/'plan/plan.json'


def test_local_validation_plan_has_no_calls_and_execution_is_bounded(tmp_path,monkeypatch):
    module,corpus,plan=local_plan(tmp_path,monkeypatch)
    prepared=json.loads(plan.read_text())
    assert len(prepared['samples'])==3 and prepared['automatic_retries']==0 and prepared['execute'] is False
    with pytest.raises(ValueError,match='allow-external'):module.execute(plan,corpus,tmp_path/'unauthorized')
    calls=[]
    class Client:
        failure_count=0
        request_count=0
        usage={}
        failures=[]
        response_models={'unit-model'}
        def __init__(self,**kwargs):assert kwargs['max_output_tokens']==2048 and kwargs['allow_fallback'] is False
        def extract(self,gene,disease,c):
            calls.append((gene,disease,c));self.request_count+=1
            return [ExtractionResult(evidence_sentence=c['context'],expression_direction='increased')]
    monkeypatch.setattr(module,'LLMClient',Client)
    report=module.execute(plan,corpus,tmp_path/'executed',allow_external=True)
    assert len(calls)==report['actual_requests']==3 and all(len(c[2]['context'])<=3000 for c in calls)
    assert report['scientific_accuracy_measured'] is False and (tmp_path/'executed/evidence.xlsx').is_file()
    with pytest.raises(ValueError,match='已存在'):module.execute(plan,corpus,tmp_path/'executed',allow_external=True)
    assert len(calls)==3


def test_local_validation_rejects_changed_plan_before_client(tmp_path,monkeypatch):
    module,corpus,plan=local_plan(tmp_path,monkeypatch)
    prepared=json.loads(plan.read_text());prepared['samples'][0]['chunk']['context']='Changed text'
    plan.write_text(json.dumps(prepared))
    monkeypatch.setattr(module,'LLMClient',lambda **kwargs:pytest.fail('Client must not be constructed'))
    with pytest.raises(ValueError,match='改变'):module.execute(plan,corpus,tmp_path/'rejected',allow_external=True)
    assert not (tmp_path/'rejected').exists()


def test_local_validation_interruption_cannot_be_replayed(tmp_path,monkeypatch):
    module,corpus,plan=local_plan(tmp_path,monkeypatch)
    class InterruptedClient:
        failure_count=0;request_count=0;usage={}
        def __init__(self,**kwargs):pass
        def extract(self,*args):self.request_count+=1;raise RuntimeError('simulated interruption')
    monkeypatch.setattr(module,'LLMClient',InterruptedClient)
    output=tmp_path/'interrupted'
    with pytest.raises(RuntimeError):module.execute(plan,corpus,output,allow_external=True)
    progress=json.loads((output/'progress.json').read_text())
    assert progress['stage']=='before_request' and progress['accounting_complete_for_started_requests'] is False
    with pytest.raises(ValueError,match='已存在'):module.execute(plan,corpus,output,allow_external=True)


def test_reference_filter_keeps_research_text_and_can_be_disabled(tmp_path):
    from source_quality import source_flags
    reference=' '.join(f'{i}. Author AB. IL6 in rheumatoid arthritis. Journal (2020). doi: 10.1000/item{i}' for i in range(1,7))
    references=chunk('reference-paper',1,reference)
    assert 'reference_like' in source_flags(references)
    research=chunk('study',1,'IL6 expression was elevated in rheumatoid arthritis patients. doi: 10.1000/article')
    assert 'reference_like' not in source_flags(research)
    corpus=tmp_path/'corpus.sqlite';CorpusStore.build(corpus,[references,research])
    store=CorpusStore(corpus)
    assert {c['paper_id'] for c in store.search('IL6','rheumatoid arthritis',10)}=={'reference-paper','study'}
    assert [c['paper_id'] for c in store.search('IL6','rheumatoid arthritis',10,skip_reference_like=True)]==['study']
    assert [p['paper_id'] for p in store.search_papers('IL6','rheumatoid arthritis',10,skip_reference_like=True)]==['study']


def test_source_hints_are_inherited_from_text_instead_of_model():
    from validators import ground_result
    c=chunk('review',1);c['title']='A systematic review of evidence'
    row=ground_result({'evidence_sentence':c['context'],'source_quality_flags':[]},c,'IL6','rheumatoid arthritis')
    assert row['source_quality_flags']==['review_title']
