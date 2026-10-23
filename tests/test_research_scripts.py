import copy
import importlib.util
import json
import hashlib
from pathlib import Path
import pytest
from openpyxl import load_workbook
from test_evaluation import case


def script(name):
    spec=importlib.util.spec_from_file_location(name,Path(__file__).parents[1]/'scripts'/f'{name}.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    return module


def pilot_case(gene='IL6'):
    c=case(gene,reviewed=False)
    c.update(gene=gene,disease='rheumatoid arthritis',title='Test abstract',source_url='https://example.org')
    return c


def test_workbook_requires_explicit_human_decisions_and_preserves_source(tmp_path):
    module=script('annotation_sheet');c=pilot_case();book=tmp_path/'annotations.xlsx'
    module.export_sheet([c],book)
    assert module.import_sheet([c],book,tmp_path/'unreviewed.jsonl')==0
    with pytest.raises(ValueError): module.export_sheet([c],book)
    workbook=load_workbook(book);sheet=workbook['案例']
    sheet['H2']='已确认';sheet['I2']='不相关';sheet['K2']='expert'
    workbook.save(book)
    with pytest.raises(ValueError,match='空关系'): module.import_sheet([c],book,tmp_path/'gold.jsonl')
    sheet['J2']='是';workbook.save(book)
    assert module.import_sheet([c],book,tmp_path/'gold.jsonl')==1
    result=json.loads((tmp_path/'gold.jsonl').read_text())
    assert result['gold_relations']==[] and result['gold_relevant'] is False
    sheet['G2']='Modified original text';workbook.save(book)
    with pytest.raises(ValueError,match='原文'): module.import_sheet([c],book,tmp_path/'tampered.jsonl')


def test_workbook_imports_positive_relation(tmp_path):
    module=script('annotation_sheet');c=pilot_case();book=tmp_path/'positive.xlsx'
    module.export_sheet([c],book);workbook=load_workbook(book)
    sheet=workbook['案例'];sheet['H2']='已确认';sheet['I2']='相关';sheet['K2']='expert'
    relation=case()['gold_relations'][0]
    for key in module.RELATION_FIELDS: workbook['关系'].cell(2,module.RELATION_COLUMNS.index(key)+1,relation.get(key,''))
    workbook.save(book)
    assert module.import_sheet([c],book,tmp_path/'positive.jsonl')==1
    assert json.loads((tmp_path/'positive.jsonl').read_text())['gold_relations'][0]['evidence_sentence']==c['context']


def test_reexport_retains_reviewed_labels_and_explicit_unreviewing_clears_them(tmp_path):
    module=script('annotation_sheet');c=pilot_case();gold=case()
    for key in ['gold_status','gold_relevant','gold_relations','reviewer']: c[key]=copy.deepcopy(gold[key])
    book=tmp_path/'reviewed.xlsx';module.export_sheet([c],book)
    assert module.import_sheet([c],book,tmp_path/'retained.jsonl')==1
    workbook=load_workbook(book);workbook['案例']['H2']='待标注';workbook.save(book)
    assert module.import_sheet([c],book,tmp_path/'cleared.jsonl')==0
    assert json.loads((tmp_path/'cleared.jsonl').read_text())['gold_relations'] is None


@pytest.mark.parametrize('field,replacement',[('gene','ABCC8'),('disease','type 2 diabetes')])
def test_relation_target_context_is_visible_and_cannot_be_swapped(tmp_path,field,replacement):
    module=script('annotation_sheet');c=pilot_case();book=tmp_path/'targets.xlsx'
    module.export_sheet([c],book);workbook=load_workbook(book);sheet=workbook['关系']
    assert sheet['B2'].value==c['gene'] and sheet['C2'].value==c['disease']
    sheet.cell(2,module.RELATION_COLUMNS.index(field)+1,replacement);workbook.save(book)
    with pytest.raises(ValueError,match='基因或疾病'): module.import_sheet([c],book,tmp_path/'blocked.jsonl')
    assert not (tmp_path/'blocked.jsonl').exists()


def test_old_workbooks_still_import_without_relation_target_columns(tmp_path):
    module=script('annotation_sheet');c=pilot_case();gold=case()
    for key in ['gold_status','gold_relevant','gold_relations','reviewer']: c[key]=copy.deepcopy(gold[key])
    book=tmp_path/'old.xlsx';module.export_sheet([c],book);workbook=load_workbook(book)
    workbook['关系'].delete_cols(2,2);workbook.save(book)
    assert module.import_sheet([c],book,tmp_path/'retained.jsonl')==1
    imported=json.loads((tmp_path/'retained.jsonl').read_text())['gold_relations'][0]
    assert imported=={key:c['gold_relations'][0].get(key,'') for key in module.RELATION_FIELDS}


def test_extended_reference_workbook_roundtrips_and_preserves_target_identity(tmp_path):
    from evaluation import evaluate
    from test_evaluation import scoped_case,prediction
    module=script('annotation_sheet');c=pilot_case();gold=scoped_case()
    for key in ['gold_status','gold_relevant','gold_relations','reviewer']:c[key]=copy.deepcopy(gold[key])
    book=tmp_path/'extended.xlsx';module.export_sheet([c],book,profile='scope_measurement')
    assert module.import_sheet([c],book,tmp_path/'extended.jsonl')==1
    imported=json.loads((tmp_path/'extended.jsonl').read_text())
    assert imported['gold_relations'][0]['measurement_type']=='unclear'
    assert evaluate([imported],[prediction(imported)],profile='scope_measurement')['metrics']['true_positive_relations']==1
    wb=load_workbook(book);wb['关系']['B2']='OTHER_GENE';wb.save(book)
    with pytest.raises(ValueError,match='基因或疾病'):module.import_sheet([c],book,tmp_path/'wrong-target.jsonl')


def test_extended_reference_does_not_turn_empty_measurement_into_unclear(tmp_path):
    module=script('annotation_sheet');c=pilot_case();book=tmp_path/'blank-extended.xlsx'
    module.export_sheet([c],book,profile='scope_measurement');wb=load_workbook(book)
    wb['案例']['H2']='已确认';wb['案例']['I2']='相关';wb['案例']['K2']='expert'
    relation=case()['gold_relations'][0]
    for k,v in relation.items():wb['关系'].cell(2,module.EXTENDED_RELATION_COLUMNS.index(k)+1,v)
    wb.save(book)
    with pytest.raises(ValueError,match='scope or measurement'):module.import_sheet([c],book,tmp_path/'not-complete.jsonl')
    assert not (tmp_path/'not-complete.jsonl').exists()


def test_evaluation_cli_selects_seven_field_profile(tmp_path,monkeypatch):
    import runpy,sys
    from test_evaluation import scoped_case,prediction
    c=scoped_case();p=prediction(c);p['relations'][0]['population_context']='healthy_only'
    annotations=tmp_path/'labels.jsonl';annotations.write_text(json.dumps(c))
    predictions=tmp_path/'predictions.jsonl';predictions.write_text(json.dumps(p))
    output=tmp_path/'score.json'
    monkeypatch.setattr(sys,'argv',['evaluate_pilot','--annotations',str(annotations),'--predictions',str(predictions),
        '--profile','scope_measurement','--output',str(output)])
    runpy.run_path(str(Path(__file__).parents[1]/'scripts/evaluate_pilot.py'),run_name='__main__')
    report=json.loads(output.read_text())
    assert report['evaluation_profile']=='scope_measurement' and report['metrics']['f1']==0


def test_chinese_screening_exposes_target_and_source_but_does_not_assign_answers(tmp_path):
    module=script('annotation_sheet');c=pilot_case();c['title']='=1+1';book=tmp_path/'simple.xlsx'
    module.export_screening_sheet([c],book);original=book.read_bytes();workbook=load_workbook(book)
    assert workbook.active.title=='只填这页'
    assert workbook['只填这页']['B2'].value=='IL6' and workbook['只填这页']['C2'].value=='类风湿关节炎'
    assert all(workbook['只填这页'].cell(2,col).value is None for col in range(5,11))
    assert workbook['原文与来源']['F2'].value==c['context']
    assert workbook['原文与来源']['D2'].data_type=='s'
    choices=[d.formula1 for d in workbook['只填这页'].data_validations.dataValidation]
    assert '"有,没有,摘要信息不足"' in choices
    with pytest.raises(ValueError,match='完整参考答案'): module.import_sheet([c],book,tmp_path/'not-gold.jsonl')
    with pytest.raises(ValueError,match='覆盖'): module.export_screening_sheet([c],book)
    assert book.read_bytes()==original and not (tmp_path/'not-gold.jsonl').exists()


def test_chinese_screening_refuses_silent_source_truncation(tmp_path):
    module=script('annotation_sheet');c=pilot_case();c['context']='x'*32768
    c['context_sha256']=hashlib.sha256(c['context'].encode()).hexdigest()
    with pytest.raises(ValueError,match='截断'): module.export_screening_sheet([c],tmp_path/'too-long.xlsx')
    assert not (tmp_path/'too-long.xlsx').exists()


def test_linked_model_config_does_not_import_unrelated_settings(tmp_path,monkeypatch):
    import config
    linked=tmp_path/'original.env'
    linked.write_text('LLM_API_KEY=private-test-value\nLLM_MODEL=linked-model\nLLM_BASE_URL=https://example.org/v1\nRAG_ENABLE_LIVE=true\n')
    monkeypatch.setattr(config,'PROJECT_ROOT',tmp_path)
    monkeypatch.setenv('RAG_MODEL_ENV_FILE',str(linked))
    for key in ['LLM_API_KEY','LLM_MODEL','LLM_BASE_URL','RAG_ENABLE_LIVE']:
        monkeypatch.delenv(key,raising=False)
    settings=config.get_llm_settings()
    assert settings['model']=='linked-model' and settings['api_key']=='private-test-value'
    import os
    assert os.getenv('RAG_ENABLE_LIVE') is None and os.getenv('LLM_API_KEY') is None
    monkeypatch.setenv('LLM_MODEL','local-model')
    assert config.get_llm_settings()['model']=='local-model'


def test_pilot_is_deterministic_and_never_creates_gold(tmp_path):
    module=script('prepare_pilot');root=tmp_path/'cache';(root/'pubmed').mkdir(parents=True)
    key='a'*32
    (root/'manifest.json').write_text(json.dumps({key:dict(gene='IL6',disease='rheumatoid arthritis')}))
    rows=[dict(chunk_id=f'p{i}',paper_id=f'PMID{i}',title='Test',context=f'IL6 expression in rheumatoid arthritis case {i}.') for i in range(3)]
    (root/'pubmed'/f'{key}.jsonl').write_text('\n'.join(json.dumps(r) for r in rows))
    first=tmp_path/'first';second=tmp_path/'second'
    module.prepare(root,first,per_query=3);module.prepare(root,second,per_query=3)
    assert (first/'annotations.jsonl').read_bytes()==(second/'annotations.jsonl').read_bytes()
    cases=[json.loads(r) for r in (first/'annotations.jsonl').read_text().splitlines()]
    assert all(c['gold_status']=='unreviewed' and c['gold_relations'] is None for c in cases)
    with pytest.raises(ValueError): module.prepare(root,first,per_query=3)


def test_model_validation_defaults_to_no_network_and_requires_both_flags(tmp_path,monkeypatch):
    module=script('validate_model');pilot=tmp_path/'pilot';pilot.mkdir()
    (pilot/'annotations.jsonl').write_text(json.dumps(pilot_case()))
    monkeypatch.setattr(module,'get_llm_settings',lambda:dict(api_key='private-test-value',model='test',base_url='https://example.org'))
    def forbidden(**kwargs): raise AssertionError('A default plan must not instantiate the model client')
    monkeypatch.setattr(module,'LLMClient',forbidden)
    report=module.run(pilot,tmp_path/'plan')
    assert report['execute'] is False and report['automatic_retries']==0
    assert 'private-test-value' not in (tmp_path/'plan/plan.json').read_text()
    with pytest.raises(ValueError,match='allow-external'): module.run(pilot,tmp_path/'blocked',execute=True)
    with pytest.raises(ValueError): module.run(pilot,tmp_path/'too-many',max_requests=4)


def test_model_validation_budget_clipping_and_repeat_guard(tmp_path,monkeypatch):
    module=script('validate_model');pilot=tmp_path/'pilot';pilot.mkdir()
    cases=[pilot_case(g) for g in ['IL6','ABCC8','CD36']]
    cases[0]['context']+='x'*4000
    cases[0]['context_sha256']=hashlib.sha256(cases[0]['context'].encode()).hexdigest()
    (pilot/'annotations.jsonl').write_text('\n'.join(json.dumps(c) for c in cases))
    monkeypatch.setattr(module,'get_llm_settings',lambda:dict(api_key='test',model='test',base_url='https://example.org'))
    calls=[]
    class MockClient:
        failure_count=0;request_count=0;usage={};failures=[]
        def __init__(self,**kwargs):
            assert kwargs['allow_fallback'] is False and kwargs['max_output_tokens']==1024
        def extract(self,gene,disease,chunk):
            calls.append(chunk);self.request_count+=1
            assert len(chunk['context'])<=3000
            return []
    monkeypatch.setattr(module,'LLMClient',MockClient)
    output=tmp_path/'execution'
    result=module.run(pilot,output,execute=True,allow_external=True)
    assert result['actual_requests']==3 and len(calls)==3
    with pytest.raises(FileExistsError): module.run(pilot,output,execute=True,allow_external=True)
    assert len(calls)==3


def test_model_recheck_plan_exposes_exact_input_and_configurable_budget_without_calling(tmp_path,monkeypatch):
    module=script('validate_model');pilot=tmp_path/'pilot';pilot.mkdir()
    c=pilot_case();c['context']+='x'*4000
    c['context_sha256']=hashlib.sha256(c['context'].encode()).hexdigest()
    (pilot/'annotations.jsonl').write_text(json.dumps(c))
    monkeypatch.setattr(module,'get_llm_settings',lambda:dict(api_key='private-test-value',model='test',base_url='https://example.org'))
    def forbidden(**kwargs): raise AssertionError('Planning must not construct a model client')
    monkeypatch.setattr(module,'LLMClient',forbidden)
    report=module.run(pilot,tmp_path/'plan',max_output_tokens=2048)
    assert report['execute'] is False and report['automatic_retries']==0
    assert report['max_output_tokens_per_request']==report['max_total_output_tokens']==2048
    sample=report['samples'][0]
    assert sample['context_sha256']==c['context_sha256'] and sample['truncated_input']
    assert sample['sent_context_characters']==3000
    assert sample['sent_context_sha256']==hashlib.sha256(c['context'][:3000].encode()).hexdigest()
    assert 'private-test-value' not in (tmp_path/'plan/plan.json').read_text()


@pytest.mark.parametrize('limit',[0,2049])
def test_validation_rejects_output_over_budget_before_model_or_output(tmp_path,monkeypatch,limit):
    module=script('validate_model')
    def forbidden(**kwargs): raise AssertionError('Invalid budgets must not construct a model client')
    monkeypatch.setattr(module,'LLMClient',forbidden)
    with pytest.raises(ValueError,match='1..2048'):
        module.run(tmp_path/'unused',tmp_path/'output',execute=True,allow_external=True,max_output_tokens=limit)
    assert not (tmp_path/'output').exists()


def test_model_validation_execution_uses_the_reviewed_output_budget(tmp_path,monkeypatch):
    module=script('validate_model');pilot=tmp_path/'pilot';pilot.mkdir()
    (pilot/'annotations.jsonl').write_text(json.dumps(pilot_case()))
    monkeypatch.setattr(module,'get_llm_settings',lambda:dict(api_key='test',model='test',base_url='https://example.org'))
    budgets=[]
    class MockClient:
        failure_count=0;request_count=0;usage={};failures=[]
        def __init__(self,**kwargs): budgets.append(kwargs['max_output_tokens'])
        def extract(self,*args): self.request_count+=1;return []
    monkeypatch.setattr(module,'LLMClient',MockClient)
    report=module.run(pilot,tmp_path/'output',execute=True,allow_external=True,max_output_tokens=2048)
    assert budgets==[2048] and report['actual_requests']==1
    assert report['max_output_tokens_per_request']==2048


def test_interrupted_validation_preserves_known_usage_and_never_replays(tmp_path,monkeypatch):
    module=script('validate_model');pilot=tmp_path/'pilot';pilot.mkdir()
    cases=[pilot_case('IL6'),pilot_case('ABCC8')]
    (pilot/'annotations.jsonl').write_text('\n'.join(json.dumps(c) for c in cases))
    monkeypatch.setattr(module,'get_llm_settings',lambda:dict(api_key='test',model='test',base_url='https://example.org'))
    calls=[]
    class InterruptingClient:
        failure_count=0;request_count=0;usage={'prompt_tokens':0,'completion_tokens':0,'total_tokens':0};failures=[]
        def __init__(self,**kwargs): pass
        def extract(self,gene,disease,chunk):
            calls.append(chunk['chunk_id']);self.request_count+=1
            if len(calls)==2: raise KeyboardInterrupt('simulated interruption while waiting for response')
            self.usage={'prompt_tokens':10,'completion_tokens':2,'total_tokens':12}
            return []
    monkeypatch.setattr(module,'LLMClient',InterruptingClient)
    output=tmp_path/'interrupted'
    with pytest.raises(KeyboardInterrupt): module.run(pilot,output,execute=True,allow_external=True)
    progress=json.loads((output/'progress.json').read_text())
    assert progress['stage']=='before_request' and progress['pending_case_id']==cases[1]['case_id']
    assert progress['completed_case_ids']==[cases[0]['case_id']]
    assert progress['actual_requests_recorded']==1 and progress['maximum_possible_requests']==2
    assert progress['usage_recorded']['total_tokens']==12
    assert not progress['accounting_complete_for_started_requests'] and not progress['execution_completed']
    with pytest.raises(FileExistsError): module.run(pilot,output,execute=True,allow_external=True)
    assert len(calls)==2
