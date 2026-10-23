import json
import pytest
from openpyxl import load_workbook
from test_research_scripts import script
from test_baseline import TEXT


def cases():
    return [dict(case_id='A',paper_id='PMID1',gene='IL6',disease='rheumatoid arthritis',split='test',
        source_url='https://pubmed.ncbi.nlm.nih.gov/1/',context=TEXT,gold_status='reviewed',gold_relations=[{'old_answer':'must not prefill'}])]


def test_new_task_sheet_does_not_reuse_legacy_answers_or_score_blanks(tmp_path):
    tool=script('baseline_reference');path=tmp_path/'new.xlsx';tool.export(cases(),path)
    book=load_workbook(path)
    assert book['1逐篇判断']['D2'].value=='待填写'
    assert book['2表达证据']['D2'].value is None
    report=tool.import_reference(cases(),path,tmp_path/'import.jsonl')
    assert report['complete_reference']==0 and report['pending']==1 and not report['accuracy_computed']
    row=json.loads((tmp_path/'import.jsonl').read_text())
    assert row['gold_relevant'] is None and row['gold_relations']==[]
    with pytest.raises(ValueError,match='覆盖'):tool.export(cases(),path)


def fill_case(book,decision):
    for col,value in [('D','已完成'),('E',decision),('F','复核者')]:book['1逐篇判断'][col+'2']=value


def test_negative_is_explicit_and_source_changes_or_formulas_are_rejected(tmp_path):
    tool=script('baseline_reference');path=tmp_path/'new.xlsx';tool.export(cases(),path)
    book=load_workbook(path);fill_case(book,'所给文本中没有符合比较');book.save(path)
    assert tool.import_reference(cases(),path,tmp_path/'negative.jsonl')['complete_reference']==1
    book['原文与来源']['G2']='Changed source.';book.save(path)
    with pytest.raises(ValueError,match='原文'):tool.import_reference(cases(),path,tmp_path/'bad.jsonl')
    book['原文与来源']['G2']=TEXT;book['1逐篇判断']['G2']='=1+1';book.save(path)
    with pytest.raises(ValueError,match='公式'):tool.import_reference(cases(),path,tmp_path/'formula.jsonl')


def test_positive_requires_both_exact_quotes_and_complete_fields(tmp_path):
    tool=script('baseline_reference');path=tmp_path/'new.xlsx';tool.export(cases(),path)
    book=load_workbook(path);fill_case(book,'有疾病与正常对照比较');book.save(path)
    with pytest.raises(ValueError,match='至少一条'):tool.import_reference(cases(),path,tmp_path/'missing.jsonl')
    values=['rheumatoid arthritis patients','healthy controls','升高',TEXT,TEXT,'转录本数量','人','blood','未报告']
    for col,value in enumerate(values,4):book['2表达证据'].cell(2,col,value)
    book.save(path)
    report=tool.import_reference(cases(),path,tmp_path/'positive.jsonl')
    row=json.loads((tmp_path/'positive.jsonl').read_text())
    assert report['complete_reference']==1 and row['gold_relations'][0]['expression_direction']=='increased'
    book['2表达证据']['H2']='A fabricated comparison of cases and controls.';book.save(path)
    with pytest.raises(ValueError,match='原句'):tool.import_reference(cases(),path,tmp_path/'fake.jsonl')
