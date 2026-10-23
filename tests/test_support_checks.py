import copy
import hashlib
import importlib.util
import json
from pathlib import Path
import pytest
from conftest import wait_for
from jobs import JobManager,Query
from web_settings import WebSettings
from validators import ground_result
from schemas import normalize_result,ExtractionResult
from support_checks import check_support


DIRECT='IL6 transcript abundance was elevated in rheumatoid arthritis patients.'


def supported():
    return dict(evidence_sentence=DIRECT,support_sentence=DIRECT,measurement_sentence=DIRECT,
        gene_or_target='IL6',disease='rheumatoid arthritis',relation_scope='direct',population_context='disease_cohort',
        measurement_type='transcript',expression_direction='increased',disease_effect='biomarker_only',confidence=.9)


def test_absence_of_flags_never_certifies_correctness():
    checks=check_support(supported(),DIRECT,'IL6','rheumatoid arthritis')
    assert checks['support_flags']==[] and checks['support_status']=='unassessed'
    assert checks['relationship_validated'] is False


def test_indirect_biomarker_quote_is_flagged_without_rewriting():
    quote='IL-6 was associated with higher marker Z levels.'
    context='We studied rheumatoid arthritis patients. '+quote
    raw=dict(evidence_sentence=quote,relation_scope='indirect',notes='This is not directly an IL6 disease relationship.',expression_direction='increased')
    chunk=dict(context=context,paper_id='PMID123',title='Test',chunk_id='c',source_type='pubmed_abstract')
    row=ground_result(raw,chunk,'IL6','rheumatoid arthritis')
    assert row['quote_verified'] and row['expression_direction']=='increased'
    assert 'indirect_or_background' in row['support_flags'] and 'model_notes_indirect' in row['support_flags']
    assert 'target_gene_missing' not in row['support_flags']
    assert 'gene_not_in_evidence' not in row['notes']
    assert row['needs_review']


def test_healthy_drug_population_is_not_silently_disease_cohort():
    context='Metformin is used for type 2 diabetes. We measured transcript changes in healthy volunteers. ABCC8 transcript levels changed after treatment.'
    row=dict(evidence_sentence='ABCC8 transcript levels changed after treatment.',population_context='disease_cohort')
    checks=check_support(row,context,'ABCC8','type 2 diabetes')
    assert 'healthy_population_hint' in checks['support_flags']
    assert row['population_context']=='disease_cohort'  # A hint must not rewrite the original model claim.


def test_healthy_controls_do_not_erase_disease_comparison():
    text='IL6 transcripts were elevated in rheumatoid arthritis patients compared with healthy individuals.'
    checks=check_support(dict(evidence_sentence=text,population_context='mixed'),text,'IL6','rheumatoid arthritis')
    assert 'healthy_population_hint' not in checks['support_flags']


def test_surface_abundance_does_not_become_total_expression():
    quote='Sarcolemmal FAT/CD36 abundance increased in type 2 diabetes patients.'
    context='This was not associated with altered expression of FAT/CD36. '+quote+' FAT/CD36 trafficking was altered.'
    row=dict(evidence_sentence=quote,expression_direction='increased',measurement_type='total_protein')
    checks=check_support(row,context,'CD36','type 2 diabetes')
    assert 'surface_measurement_hint' in checks['support_flags']
    assert 'localization_expression_distinction' in checks['support_flags']
    assert checks['measurement_hint']=='surface_abundance' and row['measurement_type']=='total_protein'


@pytest.mark.parametrize('claimed_type',['transcript','total_protein'])
def test_generic_gene_expression_quote_does_not_substantiate_specific_quantity(claimed_type):
    quote='GENE_X gene expression increased after treatment in a mouse model of disease Y.'
    row=dict(evidence_sentence=quote,support_sentence=quote,measurement_sentence=quote,
        measurement_type=claimed_type,expression_direction='increased',population_context='animal_disease_model',
        relation_scope='direct',notes='The assay was probably an RNA assay.')
    before=copy.deepcopy(row)
    checks=check_support(row,quote,'GENE_X','disease Y')
    assert checks['measurement_quote_verified']
    assert 'measurement_type_unsubstantiated' in checks['support_flags']
    assert checks['relationship_validated'] is False and row==before


@pytest.mark.parametrize(('measurement_type','quote'),[
    ('transcript','GENE_X mRNA levels increased in disease Y cells.'),
    ('transcript','GENE_X transcript abundance increased in disease Y cells.'),
    ('transcript','GENE_X mRNA rose.'),
    ('total_protein','GENE_X protein levels increased in disease Y cells.'),
    ('circulating_protein','Serum GENE_X protein concentrations increased in disease Y patients.'),
    ('surface_abundance','GENE_X cell-surface abundance increased in disease Y cells.'),
    ('localization','GENE_X translocation increased in disease Y cells.'),
    ('activity','GENE_X enzymatic activity increased in disease Y cells.'),
])
def test_explicit_target_measurement_cues_remain_review_hints(measurement_type,quote):
    row=dict(evidence_sentence=quote,support_sentence=quote,measurement_sentence=quote,
        measurement_type=measurement_type,relation_scope='direct',population_context='cell_disease_model')
    checks=check_support(row,quote,'GENE_X','disease Y')
    assert 'measurement_type_unsubstantiated' not in checks['support_flags']
    assert checks['relationship_validated'] is False


@pytest.mark.parametrize('context_prefix',[
    'Marker Z mRNA levels were measured. ',
    'RNA was isolated for analysis of a different target. ',
    '',
])
def test_other_measurements_or_pcr_do_not_substantiate_target_transcript(context_prefix):
    quote='GENE_X gene expression increased in disease Y cells after treatment, as assessed by PCR.'
    context=context_prefix+quote
    row=dict(evidence_sentence=quote,support_sentence=context,measurement_sentence=context,
        measurement_type='transcript',relation_scope='direct',population_context='cell_disease_model')
    checks=check_support(row,context,'GENE_X','disease Y')
    assert checks['measurement_quote_verified']
    assert 'measurement_type_unsubstantiated' in checks['support_flags']


def test_cell_model_observation_survives_unspecified_measurement_without_patient_extrapolation():
    scope='We used a cultured cell model of disease Y.'
    finding='Increasing regulator R raised GENE_X expression compared with control cells.'
    context=scope+' '+finding
    row=ground_result(dict(evidence_sentence=finding,support_sentence=context,measurement_sentence=finding,
        measurement_type='unclear',expression_direction='increased',population_context='cell_disease_model',
        relation_scope='direct',relation_type='associated',disease_effect='biomarker_only',species='unclear'),
        dict(context=context,paper_id='PMID123',chunk_id='c',source_type='pubmed_abstract'),'GENE_X','disease Y')
    assert row is not None and row['quote_verified'] and row['support_quote_verified']
    assert row['expression_direction']=='increased' and row['measurement_type']=='unclear'
    assert row['population_context']=='cell_disease_model' and row['species']=='unclear'
    assert row['relation_type']=='associated' and row['disease_effect']=='biomarker_only'
    assert 'measurement_unclear' in row['support_flags'] and row['needs_review']
    assert row['relationship_validated'] is False


def test_negative_and_activity_outputs_are_not_automatically_certified():
    row=supported();row.update(disease_effect='no_relation',measurement_type='activity')
    checks=check_support(row,DIRECT,'IL6','rheumatoid arthritis')
    assert 'negative_endpoint_needs_review' in checks['support_flags']
    assert 'non_expression_measurement' in checks['support_flags']
    assert row['disease_effect']=='no_relation' and checks['relationship_validated'] is False


def test_support_quotes_cannot_be_fabricated_or_certify_disease():
    row=supported();row['support_sentence']='IL6 causes every disease in humans.'
    row['measurement_sentence']='Unrelated measured quantity from nowhere.'
    checks=check_support(row,DIRECT,'IL6','rheumatoid arthritis')
    assert 'support_quote_invalid' in checks['support_flags'] and not checks['support_quote_verified']
    assert 'measurement_quote_invalid' in checks['support_flags']
    wrong=check_support(row,DIRECT,'IL6','breast cancer')
    assert 'target_disease_missing' in wrong['support_flags']


@pytest.mark.parametrize(('support','expected_flag'),[
    ('We studied rheumatoid arthritis patients.','support_gene_not_explicit'),
    ('IL6 transcript abundance increased in cultured cells.','support_disease_not_explicit'),
])
def test_support_passage_gaps_are_not_hidden_by_a_complete_expression_quote(support,expected_flag):
    row=supported();row['support_sentence']=support
    before=copy.deepcopy(row)
    context=DIRECT+' '+support
    checks=check_support(row,context,'IL6','rheumatoid arthritis')
    assert checks['support_quote_verified'] and expected_flag in checks['support_flags']
    assert row==before and not checks['relationship_validated']


def test_support_coverage_recognizes_supplied_aliases_without_certifying_relation():
    quote='IL-6 transcript abundance increased in RA patient-derived cells.'
    row=supported();row.update(evidence_sentence=quote,support_sentence=quote,measurement_sentence=quote)
    checks=check_support(row,quote,'IL6','rheumatoid arthritis',['IL-6'],['RA'])
    assert not {'support_gene_not_explicit','support_disease_not_explicit'} & set(checks['support_flags'])
    assert not checks['relationship_validated']


@pytest.mark.parametrize('invalid',[[],{},None,True])
def test_bad_new_fields_normalize_safely(invalid):
    row=normalize_result(dict(relation_scope=invalid,population_context=invalid,measurement_type=invalid,support_sentence=invalid))
    assert row.relation_scope==row.population_context==row.measurement_type=='unclear'
    assert isinstance(row.support_sentence,str)


def test_multiple_measurements_and_human_revision_remain_distinct(tmp_path,monkeypatch,owner):
    import jobs,config
    from corpus_store import CorpusStore
    quote='In rheumatoid arthritis patients, total IL6 was unchanged while surface IL6 increased.'
    db=tmp_path/'corpus.sqlite';CorpusStore.build(db,[dict(chunk_id='c',paper_id='PMID123',context=quote,title='Test')])
    monkeypatch.setattr(config,'get_llm_settings',lambda:dict(api_key='test',model='test'))
    class Client:
        failure_count=0;request_count=1
        def __init__(self,**kwargs): pass
        def extract_batch(self,gene,disease,chunks,**kwargs):
            return [(c,[ExtractionResult(evidence_sentence=quote,expression_direction='unchanged',measurement_type='total_protein'),
                ExtractionResult(evidence_sentence=quote,expression_direction='increased',measurement_type='surface_abundance')]) for c in chunks]
    monkeypatch.setattr(jobs,'LLMClient',Client)
    manager=JobManager(WebSettings(runtime_dir=tmp_path/'runtime',corpus_db=db,enable_live=True,access_password='long-test-password'))
    try:
        task=manager.submit(owner,Query(('IL6',),'rheumatoid arthritis',source='corpus',consent=True),authorized=True)
        rows=wait_for(manager,owner,task)['rows']
        assert len(rows)==2 and rows[0]['evidence_id']!=rows[1]['evidence_id']
        first=rows[0]
        manager.review(owner,task,first['evidence_id'],'corrected',changes={'measurement_type':'transcript','population_context':'mixed'},reason='Check original assay',reviewer='expert')
        changed=manager.snapshot(owner,task)['rows'][0]
        assert changed['measurement_type']=='transcript' and changed['machine_result']['measurement_type']=='total_protein'
        assert b'measurement_type' in manager.download(owner,task,'evidence.csv')
        with pytest.raises(ValueError): manager.review(owner,task,first['evidence_id'],'corrected',changes={'measurement_type':'made_up'},reason='bad',reviewer='x')
    finally: manager.close()


def test_offline_recheck_keeps_original_predictions_and_never_calls_model(tmp_path,monkeypatch):
    spec=importlib.util.spec_from_file_location('audit_predictions',Path(__file__).parents[1]/'scripts/audit_predictions.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    annotations=tmp_path/'annotations.jsonl';predictions=tmp_path/'predictions.jsonl'
    case=dict(case_id='c',paper_id='PMID123',gene='IL6',disease='rheumatoid arthritis',context=DIRECT)
    sha=hashlib.sha256(DIRECT.encode()).hexdigest()
    prediction=dict(case_id='c',paper_id='PMID123',context_sha256=sha,sent_context_sha256=sha,
        truncated_input=False,complete=True,relations=[supported()])
    annotations.write_text(json.dumps(case));predictions.write_text(json.dumps(prediction))
    before=predictions.read_bytes()
    report=module.audit(annotations,predictions,tmp_path/'audit')
    assert report['new_model_requests']==0 and report['medical_accuracy'] is None
    assert predictions.read_bytes()==before
    from openpyxl import load_workbook
    book=load_workbook(tmp_path/'audit/review-workbook.xlsx')
    assert book['待人工核对'].cell(2,13).value is None
    with pytest.raises(ValueError): module.audit(annotations,predictions,tmp_path/'audit')
    prediction['sent_context_sha256']='wrong';predictions.write_text(json.dumps(prediction))
    with pytest.raises(ValueError): module.audit(annotations,predictions,tmp_path/'bad')


def test_prompt_changes_invalidate_cache_and_state_population_measurement_requirements():
    from prompts import build_user_prompt
    from llm_cache import make_cache_key
    prompt=build_user_prompt('IL6','rheumatoid arthritis','Some text')
    assert 'healthy volunteers' in prompt and 'surface_abundance' in prompt and 'support_sentence' in prompt
    import prompts
    args=dict(model='test',base_url='https://invalid.test',gene='IL6',disease='arthritis',chunk=dict(context=DIRECT,chunk_id='c'))
    previous=prompts.PROMPT_VERSION
    try:
        current=make_cache_key(**args);prompts.PROMPT_VERSION='prior-version'
        assert current!=make_cache_key(**args)
    finally: prompts.PROMPT_VERSION=previous
