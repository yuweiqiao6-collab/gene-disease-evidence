import copy
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from retrieval import build_index
from retrieval_benchmark import (compare,digest,score,export_review,read_judgments,
    HEADERS,DECISIONS)


def bundle():
    local=[dict(chunk_id=str(i),paper_id='P'+str(i),context=f'IL6 in rheumatoid arthritis sample {i}.',
        source_type='pdf',title='Local') for i in range(4)]
    online=dict(chunk_id='online',paper_id='PMID1',context='IL6 in rheumatoid arthritis online abstract.',source_type='pubmed_abstract')
    return dict(schema='frozen-retrieval-comparison-v1',local_index=build_index(local),
        queries=[dict(query_id='Q1',gene='IL6',disease='rheumatoid arthritis',online_chunks=[online])])


def fake_ranker(calls,*,degrade=False):
    def rank(pool,gene,disease,**kw):
        calls.append(dict(pool=digest(pool),**kw))
        method='bm25'+('+dense' if kw['use_hybrid'] and not degrade else '')+('+cross_encoder' if kw['use_reranker'] else '')
        ordered=list(reversed(pool)) if kw['use_reranker'] else pool
        return ordered[:kw['top_k']],dict(actual_method=method,cache={'encoded':0,'reused':len(pool)})
    return rank


def test_all_methods_share_exact_candidates_and_online_sources_without_model_calls(tmp_path,monkeypatch):
    import retrieval_pipeline, jobs
    calls=[]
    monkeypatch.setattr(retrieval_pipeline,'rank_candidates',fake_ranker(calls))
    monkeypatch.setattr(jobs,'LLMClient',lambda **k:pytest.fail('No LLM in a ranking benchmark'))
    data=bundle();before=copy.deepcopy(data)
    report=compare(data,tmp_path/'out',ks=(1,2),repeats=2,coarse=10,encode=10)
    assert data==before and len({r['pool'] for r in calls})==1
    assert all(r['index'] is data['local_index'] for r in calls)
    assert {r['top_k'] for r in calls}=={1,2}
    assert report['quality_metrics'] is None and report['llm_calls']==0 and report['stable_rankings']
    manifest=json.loads((tmp_path/'out/review-manifest.json').read_text())
    assert len(manifest['passages'])==5 and any(r['source_type']=='pubmed_abstract' for r in manifest['passages'])
    runs=[json.loads(s) for s in (tmp_path/'out/runs.jsonl').read_text().splitlines()]
    assert len(runs)==12
    assert [r['method'] for r in runs[:6]] != [r['method'] for r in runs[6:]]
    with pytest.raises(FileExistsError):compare(data,tmp_path/'out',ks=(1,2),encode=10)


def test_fallback_is_not_reported_as_a_successful_hybrid_experiment(tmp_path,monkeypatch):
    import retrieval_pipeline
    monkeypatch.setattr(retrieval_pipeline,'rank_candidates',fake_ranker([],degrade=True))
    with pytest.raises(RuntimeError,match='降级'):compare(bundle(),tmp_path/'partial',ks=(2,),encode=10)
    assert json.loads((tmp_path/'partial/status.json').read_text())['state']=='running'
    assert not (tmp_path/'partial/report.json').exists()


def reference():
    rows=[dict(query_id='Q1',gene='IL6',disease='rheumatoid arthritis',chunk_id=str(i),review_id='R'+str(i),
        paper_id='P'+str(i),context='IL6 expression was higher in patients than in healthy controls.',title='=1+1') for i in range(3)]
    return dict(queries=[dict(query_id='Q1')],passages=rows)


def runs():
    return [dict(query_id='Q1',method=m,k=2,selected_ids=['0','1']) for m in ['bm25','hybrid','hybrid_rerank']]


def test_partial_or_uncertain_judgments_never_become_false_negatives():
    result=score(reference(),runs(),{'R0':1,'R1':None})
    assert all(r['precision_at_k'] is None and r['recall_within_frozen_pool'] is None for r in result['records'])
    assert result['insufficient']==1 and result['unfilled']==1
    partial=score(reference(),runs(),{'R0':1,'R1':0})
    assert all(r['precision_at_k']==.5 and r['recall_within_frozen_pool'] is None for r in partial['records'])
    complete=score(reference(),runs(),{'R0':1,'R1':0,'R2':1})
    assert all(r['recall_within_frozen_pool']==.5 for r in complete['records'])


def test_repeated_ranking_disagreement_cannot_be_cherry_picked_for_scoring():
    changed=runs()+[dict(query_id='Q1',method='bm25',k=2,selected_ids=['2','1'])]
    with pytest.raises(ValueError,match='排名不一致'):score(reference(),changed,{'R0':1,'R1':0,'R2':1})


def test_blind_sheet_has_blank_labels_and_rejects_changed_sources(tmp_path):
    manifest=reference();path=tmp_path/'review.xlsx'
    export_review(manifest,path)
    book=load_workbook(path);sheet=book['只填这页']
    assert [c.value for c in sheet[1]]==HEADERS
    assert all(c.value is None for row in sheet.iter_rows(min_row=2) for c in row[7:])
    assert all(sheet.cell(i,6).data_type=='s' for i in range(2,5))
    assert read_judgments(manifest,path)=={}
    sheet.cell(2,3,'another disease');book.save(path)
    with pytest.raises(ValueError,match='被修改'):read_judgments(manifest,path)


def test_positive_judgment_needs_reviewer_and_exact_source_quote(tmp_path):
    manifest=reference();path=tmp_path/'review.xlsx';export_review(manifest,path)
    book=load_workbook(path);sheet=book['只填这页'];rid=sheet.cell(2,1).value
    sheet.cell(2,8,DECISIONS[0]);sheet.cell(2,9,'not in the source');sheet.cell(2,10,'reviewer');book.save(path)
    with pytest.raises(ValueError,match='连续复制'):read_judgments(manifest,path)
    sheet.cell(2,9,manifest['passages'][0]['context']);book.save(path)
    assert read_judgments(manifest,path)=={rid:1}
    original=path.read_bytes()
    with pytest.raises(FileExistsError):export_review(manifest,path)
    assert path.read_bytes()==original


def test_over_budget_pool_is_rejected_instead_of_giving_methods_unequal_inputs(tmp_path):
    with pytest.raises(ValueError,match='共同向量预算'):compare(bundle(),tmp_path/'over',ks=(2,),encode=2)
