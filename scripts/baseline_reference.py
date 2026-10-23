"""A separate, blank Chinese reference sheet for disease-versus-normal evidence."""
import argparse
import json
import sys
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Alignment,PatternFill,Font
from openpyxl.worksheet.datavalidation import DataValidation
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from baseline import PROFILE,normalized,quote_segment
from io_utils import write_jsonl

SOURCE_KEYS=['case_id','paper_id','gene','disease','split','source_url','context']
SOURCE_HEADERS=['案例编号','论文编号','目标基因','目标疾病','原分组','来源链接','本次核对原文']
CASE_HEADERS=['案例编号','目标基因','目标疾病','完成状态','所给文本中的判断','复核人','备注']
REL_KEYS=['disease_group','normal_group','expression_direction','evidence_sentence','comparison_sentence',
          'measurement_type','species','tissue_or_cell','significance']
REL_HEADERS=['案例编号','目标基因','目标疾病','疾病组（复制原文）','正常对照（复制原文）','表达方向',
             '表达结果原句（复制原文）','比较依据原句（复制原文）','测量对象','物种','组织或细胞','统计说明']
OPTIONS={
 'expression_direction':{'升高':'increased','降低':'decreased','未变化或未见显著差异':'unchanged','不明确':'unclear'},
 'measurement_type':{'转录本数量':'transcript','总蛋白数量':'total_protein','循环蛋白浓度':'circulating_protein',
                     '膜表面丰度':'surface_abundance','不明确':'unclear'},
 'species':{'人':'human','小鼠':'mouse','大鼠':'rat','不明确':'unclear'},
 'significance':{'显著':'significant','未见显著差异':'not_significant','未报告':'not_reported','不明确':'unclear'}}
DECISIONS=['有疾病与正常对照比较','所给文本中没有符合比较','信息不足，暂不能判断']


def source_rows(cases):return [[c.get(k,'') for k in SOURCE_KEYS] for c in cases]


def export(cases,output):
    if output.exists():raise ValueError('输出已存在，不能覆盖人工工作。')
    if len({c['case_id'] for c in cases})!=len(cases):raise ValueError('案例编号重复。')
    if any(len(c['context'])>32767 for c in cases):raise ValueError('原文超出表格单元格容量，不能截断。')
    book=Workbook();guide=book.active;guide.title='先看说明'
    notes=[('目标','只找疾病状态相对正常对照的表达量变化；不填写风险、机制或治疗结论。'),
      ('第一步','打开“原文与来源”，只按给出的文字判断。摘要材料只看摘要；不要混入全文信息。'),
      ('第二步','在“1逐篇判断”填判断、复核人；核对完选择已完成。没有把握可选信息不足，备注说明。'),
      ('第三步','有符合比较时，在“2表达证据”填写组别、方向、两段原句及研究条件。原文未报告的字段可明确选不明确。'),
      ('没有符合比较','确认所给文字没有疾病与正常对照结果时，选所给文本中没有符合比较，证据页只保留前三列。不能据此说基因与疾病无关。'),
      ('不收录的比较','只有治疗前后、药物与溶剂、基因干预、患者内部亚组、基因变异、活性或定位变化，不属于本轮表达量比较。'),
      ('方向基准','一律是疾病组相对正常对照。正常组更高，意味着疾病组更低。未见显著差异不能说成证明完全相同。'),
      ('多条结果','同一案例有多条证据时，复制该案例整行追加；不同物种、组织、测量或相反方向分行。'),
      ('原句','连续复制所给文字，不要翻译、改写或拼接；疾病组与正常对照名称均应出现在比较依据原句中。'),
      ('物种与组织','物种可直接填写原文名称；组织未说明时填写不明确。人体、动物与细胞结果分开记录。'),
      ('空白与信息不足','不会作为否定答案，也不会当作已经确认的不明确；未完成案例不会充当完整人工参考。'),
      ('原始材料','原文、编号、目标和分组不可修改；导入会比对。未填旧模型答案或旧人工关系，避免误用旧任务答案。'),
      ('评测纪律','若查看本组答案后调提示，应将其转为开发材料，再另留新的独立测试；原分组仅用于来源记录。')]
    for row in notes:guide.append(row)
    guide.column_dimensions['A'].width=22;guide.column_dimensions['B'].width=105
    source=book.create_sheet('原文与来源');source.append(SOURCE_HEADERS)
    for row in source_rows(cases):source.append(row)
    decisions=book.create_sheet('1逐篇判断');decisions.append(CASE_HEADERS)
    relations=book.create_sheet('2表达证据');relations.append(REL_HEADERS)
    for c in cases:
        identity=[c['case_id'],c['gene'],c['disease']]
        decisions.append(identity+['待填写','','','']);relations.append(identity+['']*len(REL_KEYS))
    def dropdown(sheet,column,choices):
        validation=DataValidation(type='list',formula1='"'+','.join(choices)+'"',allow_blank=True)
        sheet.add_data_validation(validation);validation.add(f'{column}2:{column}5000')
    dropdown(decisions,'D',['待填写','已完成']);dropdown(decisions,'E',DECISIONS)
    from openpyxl.utils import get_column_letter
    for key,values in OPTIONS.items():dropdown(relations,get_column_letter(REL_KEYS.index(key)+4),list(values))
    for sheet in book:
        sheet.freeze_panes='D2' if sheet in (decisions,relations) else 'A2'
        for row in sheet:
            for cell in row:
                cell.alignment=Alignment(wrap_text=True,vertical='top')
                if isinstance(cell.value,str):cell.data_type='s'
                if sheet in (decisions,relations):cell.fill=PatternFill('solid',fgColor='E8EDF2' if cell.column<=3 else 'FFF2CC')
            sheet.row_dimensions[row[0].row].height=70
        for cell in sheet[1]:cell.font=Font(bold=True)
    for sheet in (decisions,relations,source):
        for column in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(column)].width=27
    source.column_dimensions['G'].width=110
    for row in range(2,source.max_row+1):source.row_dimensions[row].height=300
    relations.column_dimensions['G'].width=65;relations.column_dimensions['H'].width=65
    decisions.column_dimensions['E'].width=36
    output.parent.mkdir(parents=True,exist_ok=True);book.save(output)


def import_reference(cases,workbook,output):
    if output.exists():raise ValueError('输出已存在，不能覆盖人工工作。')
    book=load_workbook(workbook,data_only=False)
    for name,headers in [('原文与来源',SOURCE_HEADERS),('1逐篇判断',CASE_HEADERS),('2表达证据',REL_HEADERS)]:
        if name not in book or [c.value for c in book[name][1]]!=headers:raise ValueError('表头或工作表不匹配。')
    if any(c.data_type=='f' for s in book for row in s for c in row):raise ValueError('人工表不能包含公式。')
    expected=[tuple(v or None for v in row) for row in source_rows(cases)]
    if list(book['原文与来源'].iter_rows(min_row=2,values_only=True))!=expected:raise ValueError('原文、来源或分组被修改。')
    identities={c['case_id']:(c['gene'],c['disease']) for c in cases};judgments={};relations={k:[] for k in identities}
    for values in book['1逐篇判断'].iter_rows(min_row=2,values_only=True):
        cid,gene,disease,status,decision,reviewer,notes=values
        if cid not in identities or identities[cid]!=(gene,disease) or cid in judgments:raise ValueError('案例目标被修改或案例重复。')
        if status not in {'待填写','已完成'} or decision not in [None,'']+DECISIONS:raise ValueError('判断选项无效。')
        if status=='已完成' and (not decision or not str(reviewer or '').strip()):raise ValueError('完成判断需要选择结论并填写复核人。')
        judgments[cid]=(status,decision,reviewer,notes)
    if set(judgments)!=set(identities):raise ValueError('逐篇判断缺少案例。')
    context={c['case_id']:c['context'] for c in cases}
    for values in book['2表达证据'].iter_rows(min_row=2,values_only=True):
        if not any(v is not None for v in values):continue
        cid,gene,disease,*fields=values
        if cid not in identities or identities[cid]!=(gene,disease):raise ValueError('证据目标或编号被修改。')
        if not any(v not in (None,'') for v in fields):continue
        if any(v in (None,'') for v in fields):raise ValueError(f'{cid}：证据行未填写完整；原文没交代时请明确填不明确。')
        item=dict(zip(REL_KEYS,fields))
        for key,options in OPTIONS.items():
            if item[key] in options:item[key]=options[item[key]]
            elif key!='species':raise ValueError('请使用中文下拉选项。')
        for key in ('evidence_sentence','comparison_sentence'):
            if quote_segment(item[key],{'context':context[cid]}) is None:raise ValueError('原句不在所给文本内，不能改写或拼接。')
        comparison=normalized(item['comparison_sentence'])
        if any(normalized(item[k]) not in comparison for k in ('disease_group','normal_group')):raise ValueError('组别名称必须来自比较原句。')
        item.update(comparison_type='disease_vs_normal',gene_or_target=gene,disease=disease,task_profile=PROFILE)
        relations[cid].append(item)
    result=[]
    for c in cases:
        cid=c['case_id'];status,decision,reviewer,notes=judgments[cid];items=relations[cid]
        if status=='已完成' and decision==DECISIONS[0] and not items:raise ValueError('确认有比较的案例需要至少一条完整证据。')
        if decision in DECISIONS[1:] and items:raise ValueError('没有或信息不足的案例不能同时提交正证据。')
        ready=status=='已完成' and decision in DECISIONS[:2]
        row={k:c.get(k,'') for k in SOURCE_KEYS}
        row.update(task_profile=PROFILE,gold_status='reviewed' if ready else 'pending',
            gold_relevant=bool(items) if ready else None,gold_relations=items if ready else [],
            baseline_judgment=decision or '',reviewer=reviewer or '',annotation_notes=notes or '',draft_relations=items if not ready else [])
        result.append(row)
    output.parent.mkdir(parents=True,exist_ok=True);write_jsonl(output,result)
    return {'total':len(result),'complete_reference':sum(c['gold_status']=='reviewed' for c in result),
            'pending':sum(c['gold_status']!='reviewed' for c in result),'task_profile':PROFILE,'accuracy_computed':False}


if __name__=='__main__':
    parser=argparse.ArgumentParser(description='疾病与正常对照：独立中文人工参考表，不沿用旧任务答案')
    parser.add_argument('action',choices=['export','import'])
    parser.add_argument('--annotations',type=Path,required=True)
    parser.add_argument('--workbook',type=Path,required=True)
    parser.add_argument('--output',type=Path)
    args=parser.parse_args();cases=[json.loads(s) for s in args.annotations.read_text().splitlines() if s.strip()]
    if args.action=='export':export(cases,args.workbook)
    elif args.output is None:parser.error('导入需要 --output，保存到新文件。')
    else:print(json.dumps(import_reference(cases,args.workbook,args.output),ensure_ascii=False))
