"""Export/import a private expert annotation workbook without inventing gold labels."""
from __future__ import annotations
import argparse
import copy
import json
import sys
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from io_utils import write_jsonl
from evaluation import validate_labels,profile_fields,SCOPE_FIELDS,SCOPE_ALLOWED

CASE_COLUMNS=['case_id','paper_id','gene','disease','split','source_url','context',
    '人工确认状态','是否相关','确认无符合范围的关系','复核人','备注']
RELATION_FIELDS=['evidence_sentence','expression_direction','disease_effect','relation_type','species','evidence_type','tissue_or_cell']
LEGACY_RELATION_COLUMNS=['case_id']+RELATION_FIELDS
RELATION_COLUMNS=['case_id','gene','disease']+RELATION_FIELDS
EXTENDED_RELATION_COLUMNS=RELATION_COLUMNS+SCOPE_FIELDS
DISEASE_LABELS={'rheumatoid arthritis':'类风湿关节炎','type 2 diabetes':'2 型糖尿病'}
SCREEN_COLUMNS=['案例编号（点开看原文）','目标基因','目标疾病','论文编号','有没有对应结果','表达变化方向',
                '复制依据原句','比较的条件或组别','备注（可不填）','复核人']


def read_cases(path):
    return [json.loads(line) for line in path.read_text().splitlines() if line.strip()]


def export_sheet(cases,output,*,profile='legacy'):
    profile_fields(profile)
    relation_fields=RELATION_FIELDS+(SCOPE_FIELDS if profile=='scope_measurement' else [])
    relation_columns=['case_id','gene','disease']+relation_fields
    if output.exists(): raise ValueError('标注表已存在，请使用新文件名，避免覆盖人工工作。')
    output.parent.mkdir(parents=True,exist_ok=True)
    workbook=Workbook()
    sheet=workbook.active;sheet.title='案例'
    sheet.append(CASE_COLUMNS)
    for c in cases:
        reviewed=c.get('gold_status')=='reviewed'
        sheet.append([c.get(k,'') for k in CASE_COLUMNS[:7]]+[
            '已确认' if reviewed else '待标注',
            ('相关' if c['gold_relevant'] else '不相关') if reviewed else '',
            ('否' if c['gold_relations'] else '是') if reviewed else '',
            c.get('reviewer',''),c.get('annotation_notes','')])
    sheet.freeze_panes='A2'
    sheet.column_dimensions['G'].width=100
    sheet.column_dimensions['F'].width=50
    for column,options in [('H','待标注,已确认'),('I','相关,不相关'),('J','是,否')]:
        validation=DataValidation(type='list',formula1='"'+options+'"',allow_blank=True)
        sheet.add_data_validation(validation);validation.add(f'{column}2:{column}{len(cases)+1}')
    relations=workbook.create_sheet('关系')
    relations.append(relation_columns)
    for c in cases:
        if c.get('gold_status')=='reviewed' and c.get('gold_relations'):
            for r in c['gold_relations']: relations.append([c['case_id'],c['gene'],c['disease']]+[r.get(k,'') for k in relation_fields])
        else: relations.append([c['case_id'],c['gene'],c['disease']]+['']*len(relation_fields))
    relations.column_dimensions['B'].width=14
    relations.column_dimensions['C'].width=30
    relations.column_dimensions['D'].width=100
    relations.freeze_panes='D2'
    if profile=='scope_measurement':
        for field,allowed in SCOPE_ALLOWED.items():
            col=get_column_letter(relation_columns.index(field)+1)
            relations.column_dimensions[col].width=30
            validation=DataValidation(type='list',formula1='"'+','.join(sorted(allowed))+'"',allow_blank=True)
            relations.add_data_validation(validation)
            if cases: validation.add(f'{col}2:{col}{len(cases)+1}')
    help_sheet=workbook.create_sheet('填写说明')
    notes=[
        ['原则','先阅读原文，再填写参考答案；不要把模型结果直接当作标准答案。'],
        ['案例','案例编号、来源、原文和开发/测试分组不可修改。'],
        ['人工确认状态','完成检查后选择已确认，同时填写相关性与复核人。'],
        ['是否相关','当前研究问题是否存在可用的表达相关证据；不是仅仅同页提及基因和疾病。'],
        ['无关系','若没有符合任务范围的关系，在案例表明确选择确认无符合范围的关系=是。'],
        ['关系','一个案例可有多行关系；新增行保持相同 case_id。原句必须逐字来自原文。'],
        ['gene / disease（基因 / 疾病）','关系表已给出目标基因和疾病；与案例表一致，不要改成论文中的其他分子或疾病。'],
        ['方向的区别','表达方向是相对于某个对照或条件测得的升降；疾病风险或进展方向是另一问题，不能据表达升降推断治疗应抑制还是激活。'],
        ['expression_direction（表达方向）','increased=升高; decreased=降低; unchanged=未变化; unclear=不明确'],
        ['disease_effect（疾病关系）','risk_increase=风险增加; risk_decrease=风险降低; progression_increase=进展加重; progression_decrease=进展减轻; poor_prognosis=预后较差; better_prognosis=预后较好; biomarker_only=仅标志物; no_relation=未发现所述关联; unclear=不明确'],
        ['relation_type（关系类型）','associated=相关; causal=因果; mechanistic=机制; predictive=预测; therapeutic_target=治疗靶点; unclear=不明确'],
        ['species（物种）','建议 human=人; mouse=小鼠; rat=大鼠; unclear=不明确'],
        ['七字段评分（可选）','scope_measurement 模式增加以下三项；没有填写不能当作 unclear（已确认不明确）。旧表仍按旧模式使用。'],
        ['relation_scope（关系范围）','direct=直接支持目标关系（可为模型内相关性，不等于因果）; indirect=间接; background=背景; unclear=不明确'],
        ['population_context（研究对象）','disease_cohort=疾病人群; healthy_only=仅健康人; animal_disease_model=动物疾病模型; cell_disease_model=细胞疾病模型; mixed=混合; unclear=不明确'],
        ['measurement_type（测量类型）','transcript=转录本; total_protein=总蛋白; circulating_protein=循环蛋白; surface_abundance=膜表面丰度; localization=定位或转运; activity=活性; unclear=不明确'],
        ['抽样限制','仅为缓存摘要的小样本，不足以代表完整文献库；正负例和困难样本仍需人工核实。'],
        ['评测纪律','在开发组调试方法；测试组用于冻结方案后的独立检查。'],
    ]
    for row in notes: help_sheet.append(row)
    help_sheet.column_dimensions['A'].width=35;help_sheet.column_dimensions['B'].width=110
    # Force all supplied text to string cells, even if a title starts with '='.
    for sh in workbook:
        for row in sh:
            for cell in row:
                if isinstance(cell.value,str): cell.data_type='s'
    workbook.save(output)


def export_screening_sheet(cases,output):
    """Chinese first-pass review; deliberately not a complete evaluation gold sheet."""
    if output.exists(): raise ValueError('表格已存在，请使用新文件名，避免覆盖人工工作。')
    validate_labels(cases)
    if any(len(c['context'])>32767 for c in cases): raise ValueError('原文超过单元格长度限制，不能静默截断。')
    output.parent.mkdir(parents=True,exist_ok=True)
    book=Workbook();sheet=book.active;sheet.title='只填这页'
    sheet.append(SCREEN_COLUMNS)
    originals=book.create_sheet('原文与来源')
    originals.append(['案例编号','目标基因','目标疾病','论文标题','来源链接','摘要原文'])
    for index,c in enumerate(cases,2):
        disease=DISEASE_LABELS.get(c['disease'],c['disease'])
        sheet.append([c['case_id'],c['gene'],disease,c['paper_id']]+['']*6)
        sheet.cell(index,1).hyperlink=f"#'原文与来源'!A{index}"
        sheet.cell(index,1).font=Font(color='0563C1',underline='single')
        originals.append([c['case_id'],c['gene'],disease,c.get('title',''),c.get('source_url',''),c['context']])
        if c.get('source_url'): originals.cell(index,5).hyperlink=c['source_url']
    for col,options in [('E','有,没有,摘要信息不足'),('F','升高,降低,未变化,不明确,多种结果（见备注）,不适用')]:
        dropdown=DataValidation(type='list',formula1='"'+options+'"',allow_blank=True)
        dropdown.errorTitle='请使用下拉选项';dropdown.error='请选择列表中的中文选项。';dropdown.showErrorMessage=True
        sheet.add_data_validation(dropdown)
        if cases: dropdown.add(f'{col}2:{col}{len(cases)+1}')
    sheet['E1'].comment=Comment('有：原文确有当前基因、当前疾病或疾病模型中的表达/丰度研究结果。没有：只是提及、背景说明或仅基因变异。无法判断：选摘要信息不足。','填写说明')
    sheet['F1'].comment=Comment('这是表达量或丰度相对明确对照的变化，不是应该抑制还是激活该基因。有对应结果才填；多种测量方向不同，可选多种结果并写备注。','填写说明')
    sheet['G1'].comment=Comment('从原文与来源表复制连续原句，保留英文原样；不要把自己的中文总结当作论文引文。没有或信息不足时可留空。','填写说明')
    sheet['H1'].comment=Comment('例如疾病组比健康组、治疗后比治疗前。摘要未说明就写未说明，不要猜。','填写说明')
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.fill=PatternFill('solid',fgColor='E8EDF2' if cell.column<=4 else 'FFF2CC')
            cell.alignment=Alignment(vertical='top',wrap_text=True)
        sheet.row_dimensions[row[0].row].height=100
    for cell in sheet[1]:
        cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='235B6C')
        cell.alignment=Alignment(wrap_text=True,vertical='center')
    sheet.row_dimensions[1].height=35
    for col,width in {'A':32,'B':12,'C':22,'D':18,'E':23,'F':23,'G':65,'H':32,'I':40,'J':15}.items():
        sheet.column_dimensions[col].width=width
    sheet.freeze_panes='E2';sheet.auto_filter.ref=sheet.dimensions
    originals.freeze_panes='D2'
    for col,width in {'A':32,'B':12,'C':22,'D':70,'E':55,'F':115}.items(): originals.column_dimensions[col].width=width
    for row in originals:
        for cell in row: cell.alignment=Alignment(wrap_text=True,vertical='top')
    guide=book.create_sheet('怎么填')
    instructions=[
        ['只做第一轮核对','只填“只填这页”的黄色区域。灰色基因、疾病和编号已经给出，不用你找，也不要改。'],
        ['第一步：读摘要','点击案例编号跳到“原文与来源”，阅读该行原文；需要时打开来源链接。'],
        ['第二步：作判断','“有没有对应结果”选有、没有或摘要信息不足。不确定就选信息不足，不要猜。没有对应结果不等于证明没有生物学关系。'],
        ['第三步：有结果才填方向','选择升高、降低、未变化或不明确，并抄下支持的连续原句。不同测量方向不一致，选多种结果并写在备注。'],
        ['第四步：说明比较条件','例如疾病组与健康组比较，或用药后与用药前比较。原文没说明就填未说明。'],
        ['第五步：填写复核人','填写你的姓名或代号。没有或摘要信息不足时，方向和原句可以先留空。'],
        ['表达高低不等于治疗方向','基因升高而病情较重，不能直接推断抑制基因就能治疗；本表不要求你给治疗结论。'],
        ['哪些不能混为一谈','基因突变不是表达变化；健康人用药不是患者结果；动物和细胞模型请在备注说明，不能当成人体证据。'],
        ['测量要说清','转录本、总蛋白、血中蛋白浓度、膜表面丰度、定位和活性不同。原文能判断时写在备注，不能判断就留待复核。'],
        ['例子（虚构，不属于六篇论文）','原句若说某基因在疾病模型中比对照高，可填：有 / 升高 / 复制原句 / 疾病模型比对照。不要抄这个虚构例子作为真实答案。'],
        ['这不是完整评测答案','这是中文初筛意见，不自动生成专家标准答案或医学准确率。旧“关系”表这一步不用填写。'],
    ]
    for row in instructions: guide.append(row)
    guide.column_dimensions['A'].width=30;guide.column_dimensions['B'].width=115
    for row in guide:
        row[1].alignment=Alignment(wrap_text=True,vertical='top');guide.row_dimensions[row[0].row].height=55
    for sh in book:
        for row in sh:
            for cell in row:
                if isinstance(cell.value,str): cell.data_type='s'
    book.save(output)


def import_sheet(cases,workbook_path,output):
    if output.exists(): raise ValueError('输出标注已存在，请使用新文件名。')
    workbook=load_workbook(workbook_path,data_only=False)
    if '只填这页' in workbook.sheetnames:
        raise ValueError('这是中文初筛表，不是完整参考答案；请先复核初筛意见，不自动导入评测。')
    case_sheet=workbook['案例'];relation_sheet=workbook['关系']
    relation_columns=[c.value for c in relation_sheet[1]]
    if [c.value for c in case_sheet[1]]!=CASE_COLUMNS or relation_columns not in [RELATION_COLUMNS,LEGACY_RELATION_COLUMNS,EXTENDED_RELATION_COLUMNS]:
        raise ValueError('标注表列名不匹配。')
    relation_fields=RELATION_FIELDS+(SCOPE_FIELDS if relation_columns==EXTENDED_RELATION_COLUMNS else [])
    originals={c['case_id']:c for c in cases}
    annotations={};relations={key:[] for key in originals}
    for values in relation_sheet.iter_rows(min_row=2,values_only=True):
        row=dict(zip(relation_columns,values))
        if not any(v not in (None,'') for v in row.values()): continue
        case_id=row['case_id']
        if case_id not in originals: raise ValueError('关系表包含未知案例。')
        if relation_columns!=LEGACY_RELATION_COLUMNS and any(row[k]!=originals[case_id][k] for k in ['gene','disease']):
            raise ValueError('关系表中的目标基因或疾病与案例不符，不能导入。')
        if not any(row.get(k) not in (None,'') for k in relation_fields): continue
        relations[case_id].append({k:row[k] if row[k] is not None else '' for k in relation_fields})
    for values in case_sheet.iter_rows(min_row=2,values_only=True):
        row=dict(zip(CASE_COLUMNS,values))
        if not row.get('case_id'): continue
        case_id=row['case_id']
        if case_id not in originals or case_id in annotations: raise ValueError('案例编号未知或重复。')
        original=originals[case_id]
        if any(row[k]!=original.get(k,'') for k in CASE_COLUMNS[:7]): raise ValueError('原文、来源或分组被修改，不能导入。')
        case=copy.deepcopy(original)
        if row['人工确认状态']=='已确认':
            if row['是否相关'] not in {'相关','不相关'}: raise ValueError('已确认案例需要明确相关性。')
            if not relations[case_id] and row['确认无符合范围的关系']!='是':
                raise ValueError('空关系列表需要人工明确确认。')
            if relations[case_id] and row['确认无符合范围的关系']=='是':
                raise ValueError('既填写关系又确认没有关系，标注相互矛盾。')
            case.update(gold_status='reviewed',gold_relevant=row['是否相关']=='相关',
                gold_relations=relations[case_id],reviewer=str(row['复核人'] or '').strip(),annotation_notes=str(row['备注'] or ''))
        elif row['人工确认状态'] not in {'待标注',None,''}:
            raise ValueError('未知人工确认状态。')
        else:
            case.update(gold_status='unreviewed',gold_relevant=None,gold_relations=None,reviewer='',annotation_notes=str(row['备注'] or ''))
        annotations[case_id]=case
    if set(annotations)!=set(originals): raise ValueError('案例表缺少部分案例。')
    result=[annotations[c['case_id']] for c in cases]
    validate_labels(result)
    write_jsonl(output,result)
    return sum(c['gold_status']=='reviewed' for c in result)


if __name__=='__main__':
    parser=argparse.ArgumentParser(description='真实摘要的人工标注表导出或导入')
    parser.add_argument('action',choices=['export','import','screen'])
    parser.add_argument('--annotations',type=Path,required=True)
    parser.add_argument('--workbook',type=Path,required=True)
    parser.add_argument('--output',type=Path)
    parser.add_argument('--profile',choices=['legacy','scope_measurement'],default='legacy',
        help='完整表导出：legacy=原表；scope_measurement=额外三字段。导入按表头识别，中文初筛不使用此选项。')
    args=parser.parse_args()
    cases=read_cases(args.annotations)
    if args.action=='screen':
        if args.profile!='legacy': parser.error('screen is first-pass screening; use export for the seven-field reference table')
        export_screening_sheet(cases,args.workbook);print('已生成中文初筛表；只填黄色区域，不自动生成评测答案。')
    elif args.action=='export':
        export_sheet(cases,args.workbook,profile=args.profile);print('已生成标注表；没有人工确认的案例保持待标注。')
    else:
        if not args.output: parser.error('import requires --output')
        print('已导入人工确认案例数：',import_sheet(cases,args.workbook,args.output))
